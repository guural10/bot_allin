from __future__ import annotations
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

EXPORT_DIR = Path("exports")
EXPORT_DIR.mkdir(exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="5865F2")
HEADER_FONT = Font(color="FFFFFF", bold=True)
POLAND_TZ = ZoneInfo("Europe/Warsaw")

def _style(ws, widths: dict[str, int]):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

def _as_dt(value: str | None):
    """Convert an ISO Discord datetime to an Excel-safe local datetime."""
    if not value:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(POLAND_TZ).replace(tzinfo=None)
    return parsed

def export_members_xlsx(guild_name: str, records: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    headers = ["Discord ID", "Username", "Global Name", "Display Name", "Joined Server", "Account Created", "Roles", "Role Count", "Is Bot"]
    ws.append(headers)
    for r in records:
        ws.append([
            str(r["discord_id"]), r["username"], r["global_name"], r["display_name"],
            _as_dt(r["joined_at"]), _as_dt(r["account_created_at"]),
            "; ".join(role["name"] for role in r["roles"]), r["role_count"], r["is_bot"]
        ])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "yyyy-mm-dd hh:mm"
    _style(ws, {"A":22,"B":24,"C":24,"D":28,"E":20,"F":20,"G":55,"H":12,"I":10})
    if ws.max_row > 1:
        tab = Table(displayName="MembersTable", ref=f"A1:I{ws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    path = EXPORT_DIR / f"members_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    wb.save(path)
    return path

def export_roles_xlsx(roles, records: list[dict]) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "Roles"
    summary.append(["Role ID", "Role Name", "Position", "Member Count"])
    counts = {str(role.id): 0 for role in roles if not role.is_default()}
    for rec in records:
        for role in rec["roles"]:
            counts[str(role["id"])] = counts.get(str(role["id"]), 0) + 1
    for role in sorted((r for r in roles if not r.is_default()), key=lambda r: r.position, reverse=True):
        summary.append([str(role.id), role.name, role.position, counts.get(str(role.id), 0)])
    _style(summary, {"A":22,"B":35,"C":12,"D":15})

    rel = wb.create_sheet("Role Members")
    rel.append(["Discord ID", "Username", "Display Name", "Role ID", "Role Name"])
    for rec in records:
        for role in rec["roles"]:
            rel.append([str(rec["discord_id"]), rec["username"], rec["display_name"], str(role["id"]), role["name"]])
    _style(rel, {"A":22,"B":24,"C":28,"D":22,"E":35})
    path = EXPORT_DIR / f"roles_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    wb.save(path)
    return path

def export_changes_xlsx(changes: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Changes"
    ws.append(["Change Type", "Discord ID", "Username", "Display Name", "Details"])
    for r in changes["joined"]:
        ws.append(["JOINED", str(r["discord_id"]), r["username"], r["display_name"], ""])
    for r in changes["left"]:
        ws.append(["LEFT", str(r["discord_id"]), r["username"], r["display_name"], ""])
    for c in changes["role_changes"]:
        if c["added"]:
            ws.append(["ROLE_ADDED", str(c["discord_id"]), c["username"], c["display_name"], "; ".join(c["added"])])
        if c["removed"]:
            ws.append(["ROLE_REMOVED", str(c["discord_id"]), c["username"], c["display_name"], "; ".join(c["removed"])])
    _style(ws, {"A":18,"B":22,"C":24,"D":28,"E":55})
    path = EXPORT_DIR / f"changes_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    wb.save(path)
    return path


def export_access_xlsx(guild) -> Path:
    """Export effective View Channel access as a member × channel matrix plus a long relation sheet."""
    from services.access_audit import visible_channels_for_member, channel_label

    wb = Workbook()
    ws = wb.active
    ws.title = "Access Matrix"
    channels = [c for c in guild.channels if not isinstance(c, __import__("discord").CategoryChannel)]
    channels = sorted(channels, key=lambda c: (c.category.position if c.category else -1, c.position, c.name.casefold()))
    headers = ["Discord ID", "Username", "Display Name"] + [channel_label(c) for c in channels]
    ws.append(headers)
    for member in sorted(guild.members, key=lambda m: m.display_name.casefold()):
        visible_ids = {c.id for c in visible_channels_for_member(guild, member)}
        ws.append([str(member.id), member.name, member.display_name] + ["YES" if c.id in visible_ids else "" for c in channels])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 28
    for col in range(4, ws.max_column + 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 24

    rel = wb.create_sheet("Access List")
    rel.append(["Discord ID", "Username", "Display Name", "Channel ID", "Category", "Channel", "Channel Type"])
    for member in sorted(guild.members, key=lambda m: m.display_name.casefold()):
        for channel in visible_channels_for_member(guild, member):
            rel.append([str(member.id), member.name, member.display_name, str(channel.id), channel.category.name if channel.category else "", channel.name, str(channel.type)])
    _style(rel, {"A":22,"B":24,"C":28,"D":22,"E":30,"F":35,"G":18})

    path = EXPORT_DIR / f"access_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    wb.save(path)
    return path
